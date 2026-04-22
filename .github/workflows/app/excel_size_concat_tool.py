import os
import re
import traceback
from datetime import datetime
from dataclasses import dataclass
from typing import List, Optional

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


APP_TITLE = "Excel 款色尺码拼接工具"
HEADER_SCAN_MAX_ROW = 20
PREVIEW_MAX_ROWS = 12


@dataclass
class SizeColumn:
    col_idx: int
    header_text: str
    size_no: str


@dataclass
class DetectionResult:
    sheet_name: str
    header_row: int
    color_col_idx: int
    color_header: str
    size_columns: List[SizeColumn]


class ExcelTransformError(Exception):
    pass


class WorkbookAnalyzer:
    @staticmethod
    def normalize_text(value) -> str:
        if value is None:
            return ""
        text = str(value).strip()
        text = text.replace("\n", " ").replace("\r", " ")
        return text

    @staticmethod
    def extract_size_no(value) -> Optional[str]:
        if value is None:
            return None

        if isinstance(value, int):
            return str(value) if value > 0 else None

        if isinstance(value, float) and value.is_integer() and value > 0:
            return str(int(value))

        text = WorkbookAnalyzer.normalize_text(value)
        if not text:
            return None

        text_compact = text.replace(" ", "")

        m = re.fullmatch(r"尺码(\d+)", text_compact, flags=re.IGNORECASE)
        if m:
            return m.group(1)

        m = re.fullmatch(r"size(\d+)", text_compact, flags=re.IGNORECASE)
        if m:
            return m.group(1)

        if re.fullmatch(r"\d+", text_compact):
            return text_compact

        return None

    @staticmethod
    def is_color_header(value) -> bool:
        text = WorkbookAnalyzer.normalize_text(value)
        return text == "款色"

    @staticmethod
    def detect_header_row(ws) -> Optional[int]:
        best_row = None
        best_count = -1

        for row_idx in range(1, min(ws.max_row, HEADER_SCAN_MAX_ROW) + 1):
            color_found = False
            size_count = 0

            for col_idx in range(1, ws.max_column + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                if WorkbookAnalyzer.is_color_header(value):
                    color_found = True
                elif WorkbookAnalyzer.extract_size_no(value):
                    size_count += 1

            if color_found and size_count > best_count:
                best_row = row_idx
                best_count = size_count

        return best_row

    @staticmethod
    def detect_columns(ws, header_row: int) -> DetectionResult:
        color_col_idx = None
        color_header = "款色"
        size_columns: List[SizeColumn] = []

        for col_idx in range(1, ws.max_column + 1):
            value = ws.cell(row=header_row, column=col_idx).value
            text = WorkbookAnalyzer.normalize_text(value)

            if WorkbookAnalyzer.is_color_header(value):
                color_col_idx = col_idx
                color_header = text or "款色"
                continue

            size_no = WorkbookAnalyzer.extract_size_no(value)
            if size_no is not None:
                size_columns.append(
                    SizeColumn(
                        col_idx=col_idx,
                        header_text=text or f"尺码{size_no}",
                        size_no=size_no,
                    )
                )

        if color_col_idx is None:
            raise ExcelTransformError(f"在第 {header_row} 行没有找到“款色”表头。")

        if not size_columns:
            raise ExcelTransformError("在该表头行没有找到尺码列（支持：尺码1 / 1 / size1）。")

        size_columns.sort(key=lambda x: (int(x.size_no), x.col_idx))

        return DetectionResult(
            sheet_name=ws.title,
            header_row=header_row,
            color_col_idx=color_col_idx,
            color_header=color_header,
            size_columns=size_columns,
        )

    @staticmethod
    def build_preview(ws, result: DetectionResult, max_rows: int = PREVIEW_MAX_ROWS):
        preview_rows = []

        for row_idx in range(result.header_row + 1, ws.max_row + 1):
            color_value = ws.cell(row=row_idx, column=result.color_col_idx).value
            color_text = WorkbookAnalyzer.normalize_text(color_value)

            if not color_text:
                continue

            row_item = {
                "row_idx": row_idx,
                "款色": color_text,
                "values": [],
            }

            for size_col in result.size_columns:
                original_value = ws.cell(row=row_idx, column=size_col.col_idx).value
                concat_value = f"{color_text}{size_col.size_no}"
                row_item["values"].append(
                    (f"尺码{size_col.size_no}", concat_value, original_value)
                )

            preview_rows.append(row_item)

            if len(preview_rows) >= max_rows:
                break

        return preview_rows


class WorkbookTransformer:
    HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
    SUB_HEADER_FILL = PatternFill("solid", fgColor="EEF5FB")
    THIN_BORDER = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )

    @staticmethod
    def safe_sheet_name(wb, base_name: str) -> str:
        base_name = (base_name or "处理结果")[:25]
        name = base_name
        n = 1
        while name in wb.sheetnames:
            suffix = f"_{n}"
            name = f"{base_name[:31 - len(suffix)]}{suffix}"
            n += 1
        return name

    @staticmethod
    def autosize_columns(ws, max_width: int = 22):
        for column_cells in ws.columns:
            column_index = column_cells[0].column
            max_len = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_len + 2, 10), max_width
            )

    @staticmethod
    def style_result_sheet(ws, header_rows: int):
        for row in ws.iter_rows(
            min_row=1,
            max_row=header_rows,
            min_col=1,
            max_col=ws.max_column,
        ):
            for cell in row:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = WorkbookTransformer.THIN_BORDER
                cell.fill = (
                    WorkbookTransformer.HEADER_FILL
                    if cell.row == 1
                    else WorkbookTransformer.SUB_HEADER_FILL
                )

        for row in ws.iter_rows(
            min_row=header_rows + 1,
            max_row=ws.max_row,
            min_col=1,
            max_col=ws.max_column,
        ):
            for cell in row:
                cell.border = WorkbookTransformer.THIN_BORDER
                if cell.column == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.freeze_panes = f"A{header_rows + 1}"
        ws.sheet_view.showGridLines = True
        WorkbookTransformer.autosize_columns(ws)

    @staticmethod
    def create_result_sheet(wb, source_ws, detect: DetectionResult, output_mode: str, log_cb):
        sheet_name = WorkbookTransformer.safe_sheet_name(wb, f"{source_ws.title}_处理结果")
        result_ws = wb.create_sheet(sheet_name)

        if output_mode == "concat_only":
            headers = ["款色"] + [f"尺码{sc.size_no}" for sc in detect.size_columns]
            for col_idx, header in enumerate(headers, start=1):
                result_ws.cell(row=1, column=col_idx, value=header)

            target_row = 2
            for row_idx in range(detect.header_row + 1, source_ws.max_row + 1):
                color_text = WorkbookAnalyzer.normalize_text(
                    source_ws.cell(row=row_idx, column=detect.color_col_idx).value
                )
                if not color_text:
                    continue

                result_ws.cell(row=target_row, column=1, value=color_text)
                for offset, sc in enumerate(detect.size_columns, start=2):
                    result_ws.cell(row=target_row, column=offset, value=f"{color_text}{sc.size_no}")
                target_row += 1

            WorkbookTransformer.style_result_sheet(result_ws, header_rows=1)
            log_cb(f"已生成结果工作表：{sheet_name}（模式：仅输出拼接结果）")
            return sheet_name

        result_ws.cell(row=1, column=1, value="款色")
        current_col = 2

        for sc in detect.size_columns:
            result_ws.cell(row=1, column=current_col, value=f"尺码{sc.size_no}")
            result_ws.cell(row=1, column=current_col + 1, value=f"尺码{sc.size_no}")
            result_ws.cell(row=2, column=current_col, value="拼接结果")
            result_ws.cell(row=2, column=current_col + 1, value="原值")
            current_col += 2

        result_ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

        target_row = 3
        data_count = 0

        for row_idx in range(detect.header_row + 1, source_ws.max_row + 1):
            color_text = WorkbookAnalyzer.normalize_text(
                source_ws.cell(row=row_idx, column=detect.color_col_idx).value
            )
            if not color_text:
                continue

            result_ws.cell(row=target_row, column=1, value=color_text)
            current_col = 2

            for sc in detect.size_columns:
                original_value = source_ws.cell(row=row_idx, column=sc.col_idx).value
                result_ws.cell(row=target_row, column=current_col, value=f"{color_text}{sc.size_no}")
                result_ws.cell(row=target_row, column=current_col + 1, value=original_value)
                current_col += 2

            data_count += 1
            target_row += 1

        WorkbookTransformer.style_result_sheet(result_ws, header_rows=2)
        log_cb(
            f"已生成结果工作表：{sheet_name}（模式：拼接结果 + 原值并排输出），共写入 {data_count} 行数据"
        )
        return sheet_name

    @staticmethod
    def process_file(input_path: str, output_path: str, sheet_name: str, header_row: int, output_mode: str, log_cb):
        wb = load_workbook(input_path)

        if sheet_name not in wb.sheetnames:
            raise ExcelTransformError(f"工作表不存在：{sheet_name}")

        source_ws = wb[sheet_name]
        detect = WorkbookAnalyzer.detect_columns(source_ws, header_row)

        log_cb(f"开始处理工作表：{sheet_name}")
        log_cb(f"检测到 款色 列：{get_column_letter(detect.color_col_idx)}；尺码列数量：{len(detect.size_columns)}")

        result_sheet_name = WorkbookTransformer.create_result_sheet(
            wb, source_ws, detect, output_mode, log_cb
        )

        out_dir = os.path.dirname(output_path)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)

        wb.save(output_path)

        log_cb(f"文件已保存到：{output_path}")
        log_cb(f"新增工作表：{result_sheet_name}")
        return result_sheet_name


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1220x820")
        self.minsize(1080, 760)

        self.input_path_var = tk.StringVar()
        self.output_path_var = tk.StringVar()
        self.sheet_var = tk.StringVar()
        self.header_row_var = tk.IntVar(value=1)
        self.output_mode_var = tk.StringVar(value="concat_and_original")
        self.confirm_detect_var = tk.BooleanVar(value=False)
        self.confirm_preview_var = tk.BooleanVar(value=False)

        self.wb_preview = None
        self.detect_result: Optional[DetectionResult] = None
        self.preview_rows = []

        self._build_ui()

    def log(self, message: str):
        now = datetime.now().strftime("%H:%M:%S")
        self.log_text.configure(state="normal")
        self.log_text.insert("end", f"[{now}] {message}\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _build_ui(self):
        top = ttk.Frame(self, padding=12)
        top.pack(fill="x")

        file_frame = ttk.LabelFrame(top, text="文件设置", padding=10)
        file_frame.pack(fill="x")

        ttk.Label(file_frame, text="输入文件：").grid(row=0, column=0, sticky="w", pady=6)
        ttk.Entry(file_frame, textvariable=self.input_path_var).grid(row=0, column=1, sticky="ew", padx=6)
        ttk.Button(file_frame, text="选择文件", command=self.choose_input_file).grid(row=0, column=2, padx=6)
        ttk.Button(file_frame, text="加载工作簿", command=self.load_workbook_preview).grid(row=0, column=3, padx=6)

        ttk.Label(file_frame, text="输出文件：").grid(row=1, column=0, sticky="w", pady=6)
        ttk.Entry(file_frame, textvariable=self.output_path_var).grid(row=1, column=1, sticky="ew", padx=6)
        ttk.Button(file_frame, text="选择保存位置", command=self.choose_output_file).grid(row=1, column=2, padx=6)

        file_frame.columnconfigure(1, weight=1)

        config_frame = ttk.LabelFrame(top, text="识别与处理设置", padding=10)
        config_frame.pack(fill="x", pady=(10, 0))

        ttk.Label(config_frame, text="工作表：").grid(row=0, column=0, sticky="w", pady=6)
        self.sheet_combo = ttk.Combobox(config_frame, textvariable=self.sheet_var, state="readonly", width=24)
        self.sheet_combo.grid(row=0, column=1, sticky="w", padx=6)

        ttk.Label(config_frame, text="表头行：").grid(row=0, column=2, sticky="w")
        ttk.Spinbox(config_frame, from_=1, to=9999, textvariable=self.header_row_var, width=8).grid(
            row=0, column=3, sticky="w", padx=6
        )

        ttk.Label(config_frame, text="输出模式：").grid(row=0, column=4, sticky="w")
        self.mode_combo = ttk.Combobox(
            config_frame,
            textvariable=self.output_mode_var,
            state="readonly",
            width=34,
            values=[
                "concat_and_original",
                "concat_only",
            ],
        )
        self.mode_combo.grid(row=0, column=5, sticky="w", padx=6)
        self.mode_combo.set("concat_and_original")

        ttk.Label(
            config_frame,
            text="模式说明：concat_and_original=拼接结果+原值；concat_only=仅输出拼接结果",
        ).grid(row=1, column=0, columnspan=6, sticky="w", pady=(2, 6))

        btn_frame = ttk.Frame(config_frame)
        btn_frame.grid(row=2, column=0, columnspan=6, sticky="w")
        ttk.Button(btn_frame, text="自动识别表头", command=self.auto_detect).pack(side="left", padx=(0, 6))
        ttk.Button(btn_frame, text="生成预览", command=self.generate_preview).pack(side="left", padx=6)
        ttk.Button(btn_frame, text="开始处理并保存", command=self.process_and_save).pack(side="left", padx=6)

        ttk.Checkbutton(
            config_frame,
            text="我已核对识别结果",
            variable=self.confirm_detect_var,
        ).grid(row=3, column=0, columnspan=2, sticky="w", pady=(8, 0))

        ttk.Checkbutton(
            config_frame,
            text="我已核对预览结果",
            variable=self.confirm_preview_var,
        ).grid(row=3, column=2, columnspan=2, sticky="w", pady=(8, 0))

        middle = ttk.Panedwindow(self, orient="vertical")
        middle.pack(fill="both", expand=True, padx=12, pady=12)

        preview_frame = ttk.LabelFrame(middle, text="识别结果与预览", padding=10)
        middle.add(preview_frame, weight=3)

        info_frame = ttk.Frame(preview_frame)
        info_frame.pack(fill="x")

        self.detect_info = tk.Text(info_frame, height=6, wrap="word")
        self.detect_info.pack(fill="x")

        tree_frame = ttk.Frame(preview_frame)
        tree_frame.pack(fill="both", expand=True, pady=(8, 0))

        self.preview_tree = ttk.Treeview(tree_frame, columns=("row", "color", "details"), show="headings")
        self.preview_tree.heading("row", text="源数据行")
        self.preview_tree.heading("color", text="款色")
        self.preview_tree.heading("details", text="预览详情")
        self.preview_tree.column("row", width=100, anchor="center")
        self.preview_tree.column("color", width=180, anchor="w")
        self.preview_tree.column("details", width=800, anchor="w")

        yscroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.preview_tree.yview)
        self.preview_tree.configure(yscrollcommand=yscroll.set)
        self.preview_tree.pack(side="left", fill="both", expand=True)
        yscroll.pack(side="right", fill="y")

        log_frame = ttk.LabelFrame(middle, text="日志", padding=10)
        middle.add(log_frame, weight=2)

        self.log_text = tk.Text(log_frame, height=12, state="disabled")
        self.log_text.pack(fill="both", expand=True)

        self.log("工具已启动。建议流程：加载工作簿 -> 自动识别 -> 生成预览 -> 勾选两项确认 -> 开始处理。")

    def choose_input_file(self):
        path = filedialog.askopenfilename(
            title="选择 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx *.xlsm *.xltx *.xltm")],
        )
        if not path:
            return

        self.input_path_var.set(path)

        if not self.output_path_var.get().strip():
            base, ext = os.path.splitext(path)
            self.output_path_var.set(f"{base}_处理结果{ext or '.xlsx'}")

        self.log(f"已选择输入文件：{path}")

    def choose_output_file(self):
        path = filedialog.asksaveasfilename(
            title="选择保存位置",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
        )
        if not path:
            return

        self.output_path_var.set(path)
        self.log(f"已选择输出文件：{path}")

    def load_workbook_preview(self):
        input_path = self.input_path_var.get().strip()
        if not input_path:
            messagebox.showwarning(APP_TITLE, "请先选择输入文件。")
            return

        try:
            self.wb_preview = load_workbook(input_path, data_only=False)
            sheets = self.wb_preview.sheetnames
            self.sheet_combo["values"] = sheets

            if sheets:
                self.sheet_var.set(sheets[0])

            self.confirm_detect_var.set(False)
            self.confirm_preview_var.set(False)
            self.detect_result = None
            self.preview_rows = []
            self.clear_preview()

            self.log(f"工作簿加载成功，共 {len(sheets)} 个工作表：{', '.join(sheets)}")
        except Exception as e:
            messagebox.showerror(APP_TITLE, f"加载工作簿失败：{e}")
            self.log(f"加载工作簿失败：{e}")

    def get_current_ws(self):
        if self.wb_preview is None:
            raise ExcelTransformError("请先加载工作簿。")

        sheet_name = self.sheet_var.get().strip()
        if not sheet_name:
            raise ExcelTransformError("请选择工作表。")

        if sheet_name not in self.wb_preview.sheetnames:
            raise ExcelTransformError(f"工作表不存在：{sheet_name}")

        return self.wb_preview[sheet_name]

    def auto_detect(self):
        try:
            ws = self.get_current_ws()
            header_row = WorkbookAnalyzer.detect_header_row(ws)

            if not header_row:
                raise ExcelTransformError("自动识别失败：前 20 行内没有同时找到“款色”和尺码列。")

            self.header_row_var.set(header_row)
            self.detect_result = WorkbookAnalyzer.detect_columns(ws, header_row)
            self.show_detect_info(self.detect_result)
            self.confirm_detect_var.set(False)
            self.confirm_preview_var.set(False)

            self.log(f"自动识别完成：工作表={ws.title}，表头行={header_row}")
        except Exception as e:
            messagebox.showerror(APP_TITLE, str(e))
            self.log(f"自动识别失败：{e}")

    def show_detect_info(self, detect: DetectionResult):
        color_col = get_column_letter(detect.color_col_idx)
        size_desc = "，".join([f"{get_column_letter(s.col_idx)}=尺码{s.size_no}" for s in detect.size_columns])

        text = (
            f"工作表：{detect.sheet_name}\n"
            f"表头行：{detect.header_row}\n"
            f"款色列：{color_col}（{detect.color_header}）\n"
            f"识别到 {len(detect.size_columns)} 个尺码列：{size_desc}\n"
            f"提示：如果自动识别有偏差，可以手动修改“表头行”，然后重新生成预览。"
        )

        self.detect_info.delete("1.0", "end")
        self.detect_info.insert("1.0", text)

    def clear_preview(self):
        self.detect_info.delete("1.0", "end")
        for item in self.preview_tree.get_children():
            self.preview_tree.delete(item)

    def generate_preview(self):
        try:
            ws = self.get_current_ws()
            header_row = int(self.header_row_var.get())
            self.detect_result = WorkbookAnalyzer.detect_columns(ws, header_row)
            self.show_detect_info(self.detect_result)

            self.preview_rows = WorkbookAnalyzer.build_preview(ws, self.detect_result)

            for item in self.preview_tree.get_children():
                self.preview_tree.delete(item)

            for row in self.preview_rows:
                parts = []
                for size_label, concat_value, original_value in row["values"]:
                    parts.append(f"{size_label}: {concat_value} | 原值={original_value}")

                self.preview_tree.insert(
                    "",
                    "end",
                    values=(row["row_idx"], row["款色"], "；  ".join(parts)),
                )

            self.confirm_preview_var.set(False)
            self.log(f"预览已生成，共展示前 {len(self.preview_rows)} 行有效数据。")

            if not self.preview_rows:
                self.log("提示：没有可预览的有效数据行，请检查表头行或数据内容。")
        except Exception as e:
            messagebox.showerror(APP_TITLE, str(e))
            self.log(f"生成预览失败：{e}")

    def process_and_save(self):
        input_path = self.input_path_var.get().strip()
        output_path = self.output_path_var.get().strip()

        if not input_path:
            messagebox.showwarning(APP_TITLE, "请先选择输入文件。")
            return

        if not output_path:
            messagebox.showwarning(APP_TITLE, "请选择输出文件保存位置。")
            return

        if self.detect_result is None:
            messagebox.showwarning(APP_TITLE, "请先完成自动识别或生成预览。")
            return

        if not self.confirm_detect_var.get():
            messagebox.showwarning(APP_TITLE, "请先勾选“我已核对识别结果”。")
            return

        if not self.confirm_preview_var.get():
            messagebox.showwarning(APP_TITLE, "请先勾选“我已核对预览结果”。")
            return

        if os.path.abspath(input_path) == os.path.abspath(output_path):
            if not messagebox.askyesno(
                APP_TITLE,
                "输入文件和输出文件相同。是否继续？继续后会直接覆盖原文件。",
            ):
                return

        detect = self.detect_result
        size_text = "、".join([f"尺码{s.size_no}" for s in detect.size_columns])

        summary = (
            f"即将处理：\n"
            f"- 文件：{os.path.basename(input_path)}\n"
            f"- 工作表：{detect.sheet_name}\n"
            f"- 表头行：{detect.header_row}\n"
            f"- 款色列：{get_column_letter(detect.color_col_idx)}\n"
            f"- 尺码列：{size_text}\n"
            f"- 输出模式：{self.output_mode_var.get()}\n"
            f"- 保存到：{output_path}\n\n"
            f"请再次确认。"
        )

        if not messagebox.askyesno(APP_TITLE, summary):
            self.log("用户取消了第一次执行确认。")
            return

        typed = simpledialog.askstring(APP_TITLE, "数据重要，请输入：确认处理")
        if typed != "确认处理":
            self.log("最终确认未通过，已取消处理。")
            messagebox.showinfo(APP_TITLE, "未输入“确认处理”，本次未执行。")
            return

        try:
            result_sheet_name = WorkbookTransformer.process_file(
                input_path=input_path,
                output_path=output_path,
                sheet_name=detect.sheet_name,
                header_row=detect.header_row,
                output_mode=self.output_mode_var.get().strip(),
                log_cb=self.log,
            )
            messagebox.showinfo(
                APP_TITLE,
                f"处理完成。\n新增结果工作表：{result_sheet_name}\n输出文件：{output_path}",
            )
        except Exception as e:
            self.log(f"处理失败：{e}")
            self.log(traceback.format_exc())
            messagebox.showerror(APP_TITLE, f"处理失败：{e}")


if __name__ == "__main__":
    app = App()
    app.mainloop()
